"""CRUD, bulk CSV/GeoJSON upload, spatial search APIs."""

from __future__ import annotations

import csv
import io
import json
import math
import uuid
from typing import Any

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from geoalchemy2.elements import WKTElement
from sqlalchemy import func, select, text
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.api.deps import (
    assert_operator_department,
    get_current_active_user,
    require_write,
)
from app.core.config import settings
from app.db.models import Camera, CameraStatus, Department, User
from app.db.session import get_db
from app.schemas.camera import (
    BulkUploadResult,
    CameraCreate,
    CameraNearbyOut,
    CameraOut,
    CameraUpdate,
)
from app.services.spatial import find_nearby_cameras

router = APIRouter(dependencies=[Depends(get_current_active_user)])


def _fov_wkt(lon: float, lat: float, heading: float, fov: float, range_m: float) -> str:
    """Approximate FOV as a fan polygon in WGS84 (local meters → degrees)."""
    meters_per_deg_lat = 111_320.0
    meters_per_deg_lon = 111_320.0 * max(math.cos(math.radians(lat)), 0.01)
    half = fov / 2.0
    points: list[str] = [f"{lon} {lat}"]
    for i in range(9):
        angle = math.radians(heading - half + (fov * i / 8.0))
        dlon = (range_m * math.sin(angle)) / meters_per_deg_lon
        dlat = (range_m * math.cos(angle)) / meters_per_deg_lat
        points.append(f"{lon + dlon} {lat + dlat}")
    points.append(f"{lon} {lat}")
    return f"POLYGON(({' ,'.join(points)}))"


async def _get_or_create_department(db: AsyncSession, code: str | None) -> Department | None:
    if not code:
        return None
    code = code.strip().upper()
    result = await db.execute(select(Department).where(Department.code == code))
    dept = result.scalar_one_or_none()
    if dept:
        return dept
    dept = Department(code=code, name=code.replace("_", " ").title())
    db.add(dept)
    await db.flush()
    return dept


def _camera_out(cam: Camera) -> CameraOut:
    dept_code = cam.department.code if cam.department else None
    stream_id = cam.whep_path or cam.external_id
    return CameraOut(
        id=cam.id,
        external_id=cam.external_id,
        name=cam.name,
        department_code=dept_code,
        latitude=cam.latitude,
        longitude=cam.longitude,
        altitude_m=cam.altitude_m,
        heading_deg=cam.heading_deg,
        fov_deg=cam.fov_deg,
        range_m=cam.range_m,
        status=cam.status.value if isinstance(cam.status, CameraStatus) else cam.status,
        is_active=cam.is_active,
        whep_url=settings.whep_url(stream_id),
        hls_url=settings.hls_url(stream_id),
        last_seen_at=cam.last_seen_at,
        installed_at=cam.installed_at,
        meta=cam.meta,
    )


def _apply_geometry(cam: Camera) -> None:
    if cam.longitude is not None and cam.latitude is not None:
        cam.location = WKTElement(f"POINT({cam.longitude} {cam.latitude})", srid=4326)
        if cam.heading_deg is not None and cam.fov_deg and cam.range_m:
            cam.fov_polygon = WKTElement(
                _fov_wkt(cam.longitude, cam.latitude, cam.heading_deg, cam.fov_deg, cam.range_m),
                srid=4326,
            )


@router.get("", response_model=list[CameraOut])
async def list_cameras(
    department: str | None = None,
    status: str | None = None,
    limit: int = Query(500, ge=1, le=5000),
    offset: int = Query(0, ge=0),
    db: AsyncSession = Depends(get_db),
):
    q = select(Camera).options(selectinload(Camera.department)).where(Camera.is_active.is_(True))
    if department:
        q = q.join(Department).where(Department.code == department.upper())
    if status:
        q = q.where(Camera.status == CameraStatus(status))
    q = q.order_by(Camera.name).offset(offset).limit(limit)
    rows = (await db.execute(q)).scalars().all()
    return [_camera_out(c) for c in rows]


@router.get("/geojson")
async def cameras_geojson(
    department: str | None = None,
    db: AsyncSession = Depends(get_db),
):
    q = select(Camera).options(selectinload(Camera.department)).where(
        Camera.is_active.is_(True),
        Camera.latitude.is_not(None),
        Camera.longitude.is_not(None),
    )
    if department:
        q = q.join(Department).where(Department.code == department.upper())
    cameras = (await db.execute(q)).scalars().all()
    features = []
    for cam in cameras:
        features.append(
            {
                "type": "Feature",
                "geometry": {"type": "Point", "coordinates": [cam.longitude, cam.latitude]},
                "properties": {
                    "id": str(cam.id),
                    "external_id": cam.external_id,
                    "name": cam.name,
                    "department": cam.department.code if cam.department else None,
                    "status": cam.status.value if isinstance(cam.status, CameraStatus) else cam.status,
                    "heading_deg": cam.heading_deg,
                    "fov_deg": cam.fov_deg,
                    "range_m": cam.range_m,
                    "geo_source": (cam.meta or {}).get("geo_source"),
                    "latitude": cam.latitude,
                    "longitude": cam.longitude,
                },
            }
        )
    return {"type": "FeatureCollection", "features": features}


@router.get("/spatial", response_model=list[CameraOut])
async def spatial_search(
    lat: float = Query(..., ge=-90, le=90),
    lon: float = Query(..., ge=-180, le=180),
    radius_m: float = Query(1000, gt=0, le=100_000),
    department: str | None = None,
    limit: int = Query(200, ge=1, le=2000),
    db: AsyncSession = Depends(get_db),
):
    # Geography cast for metre-accurate ST_DWithin
    sql = text(
        """
        SELECT c.id
        FROM cameras c
        LEFT JOIN departments d ON d.id = c.department_id
        WHERE c.is_active = true
          AND c.location IS NOT NULL
          AND ST_DWithin(
                c.location::geography,
                ST_SetSRID(ST_MakePoint(:lon, :lat), 4326)::geography,
                :radius
              )
          AND (:dept IS NULL OR d.code = :dept)
        ORDER BY ST_Distance(
                   c.location::geography,
                   ST_SetSRID(ST_MakePoint(:lon, :lat), 4326)::geography
                 )
        LIMIT :lim
        """
    )
    result = await db.execute(
        sql,
        {"lat": lat, "lon": lon, "radius": radius_m, "dept": department.upper() if department else None, "lim": limit},
    )
    ids = [row[0] for row in result.fetchall()]
    if not ids:
        return []
    q = select(Camera).options(selectinload(Camera.department)).where(Camera.id.in_(ids))
    cams = {c.id: c for c in (await db.execute(q)).scalars().all()}
    return [_camera_out(cams[i]) for i in ids if i in cams]


@router.get("/nearby", response_model=list[CameraNearbyOut])
async def cameras_nearby(
    lat: float = Query(..., ge=-90, le=90),
    lon: float = Query(..., ge=-180, le=180),
    radius_meters: float = Query(500.0, gt=0, le=100_000),
    limit: int = Query(10, ge=1, le=100),
    department: str | None = None,
    db: AsyncSession = Depends(get_db),
    _user: User | None = Depends(get_current_active_user),
):
    """
    Tap-to-discover: cameras near a map click, sorted by distance_meters.
    Falls back to KNN nearest neighbours if the radius is empty.
    """
    rows = await find_nearby_cameras(
        db,
        lat=lat,
        lon=lon,
        radius_meters=radius_meters,
        limit=limit,
        department=department,
    )
    # Detect KNN fallback: any row beyond radius means we fell through
    used_knn = bool(rows) and all(
        float(r["distance_meters"]) > radius_meters for r in rows
    )
    out: list[CameraNearbyOut] = []
    for r in rows:
        stream_id = r.get("whep_path") or r["external_id"]
        status_raw = str(r.get("status") or "unknown").lower().replace("camerastatus.", "")
        if status_raw not in ("online", "offline", "degraded", "unknown"):
            status_raw = "unknown"
        out.append(
            CameraNearbyOut(
                id=r["id"],
                external_id=r["external_id"],
                name=r["name"],
                department_code=r.get("department_code"),
                latitude=r.get("latitude"),
                longitude=r.get("longitude"),
                altitude_m=None,
                heading_deg=r.get("heading_deg"),
                fov_deg=r.get("fov_deg"),
                range_m=r.get("range_m"),
                status=status_raw,
                is_active=True,
                whep_url=settings.whep_url(stream_id),
                hls_url=settings.hls_url(r.get("hls_path") or r["external_id"]),
                last_seen_at=None,
                installed_at=None,
                meta=r.get("meta") if isinstance(r.get("meta"), dict) else None,
                distance_meters=float(r["distance_meters"]),
                knn_fallback=used_knn or float(r["distance_meters"]) > radius_meters,
            )
        )
    return out


@router.get("/meta/departments")
async def list_departments(db: AsyncSession = Depends(get_db)):
    q = (
        select(Department, func.count(Camera.id).label("camera_count"))
        .outerjoin(Camera, Camera.department_id == Department.id)
        .group_by(Department.id)
        .order_by(Department.name)
    )
    rows = (await db.execute(q)).all()
    return [
        {"id": str(d.id), "code": d.code, "name": d.name, "camera_count": count}
        for d, count in rows
    ]


@router.get("/{camera_id}", response_model=CameraOut)
async def get_camera(camera_id: uuid.UUID, db: AsyncSession = Depends(get_db)):
    q = select(Camera).options(selectinload(Camera.department)).where(Camera.id == camera_id)
    cam = (await db.execute(q)).scalar_one_or_none()
    if not cam:
        raise HTTPException(404, "Camera not found")
    return _camera_out(cam)


@router.post("", response_model=CameraOut, status_code=201)
async def create_camera(
    payload: CameraCreate,
    db: AsyncSession = Depends(get_db),
    user: User | None = Depends(get_current_active_user),
):
    u = require_write(user)
    assert_operator_department(u, payload.department_code)
    existing = (
        await db.execute(select(Camera).where(Camera.external_id == payload.external_id))
    ).scalar_one_or_none()
    if existing:
        raise HTTPException(409, f"Camera {payload.external_id} already exists")
    dept = await _get_or_create_department(db, payload.department_code)
    cam = Camera(
        external_id=payload.external_id,
        name=payload.name,
        department_id=dept.id if dept else None,
        latitude=payload.latitude,
        longitude=payload.longitude,
        altitude_m=payload.altitude_m,
        heading_deg=payload.heading_deg,
        fov_deg=payload.fov_deg,
        range_m=payload.range_m,
        rtsp_url=payload.rtsp_url,
        whep_path=payload.whep_path or payload.external_id,
        hls_path=payload.hls_path or payload.external_id,
        installed_at=payload.installed_at,
        meta=payload.meta,
        status=CameraStatus.UNKNOWN,
    )
    _apply_geometry(cam)
    db.add(cam)
    await db.flush()
    await db.refresh(cam, attribute_names=["department"])
    if dept:
        cam.department = dept
    return _camera_out(cam)


@router.patch("/{camera_id}", response_model=CameraOut)
async def update_camera(
    camera_id: uuid.UUID,
    payload: CameraUpdate,
    db: AsyncSession = Depends(get_db),
    user: User | None = Depends(get_current_active_user),
):
    u = require_write(user)
    q = select(Camera).options(selectinload(Camera.department)).where(Camera.id == camera_id)
    cam = (await db.execute(q)).scalar_one_or_none()
    if not cam:
        raise HTTPException(404, "Camera not found")
    existing_dept = cam.department.code if cam.department else None
    assert_operator_department(u, existing_dept)
    data = payload.model_dump(exclude_unset=True)
    dept_code = data.pop("department_code", None)
    if dept_code is not None:
        assert_operator_department(u, dept_code)
        dept = await _get_or_create_department(db, dept_code)
        cam.department_id = dept.id if dept else None
        cam.department = dept
    status = data.pop("status", None)
    for k, v in data.items():
        setattr(cam, k, v)
    if status is not None:
        cam.status = CameraStatus(status.value if hasattr(status, "value") else status)
    _apply_geometry(cam)
    await db.flush()
    return _camera_out(cam)


@router.delete("/{camera_id}", status_code=204)
async def delete_camera(
    camera_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    user: User | None = Depends(get_current_active_user),
):
    u = require_write(user)
    cam = (
        await db.execute(
            select(Camera).options(selectinload(Camera.department)).where(Camera.id == camera_id)
        )
    ).scalar_one_or_none()
    if not cam:
        raise HTTPException(404, "Camera not found")
    assert_operator_department(u, cam.department.code if cam.department else None)
    cam.is_active = False


@router.post("/bulk", response_model=BulkUploadResult)
async def bulk_upload(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    user: User | None = Depends(get_current_active_user),
):
    u = require_write(user)
    raw = await file.read()
    filename = (file.filename or "").lower()
    rows: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []

    try:
        if filename.endswith(".geojson") or filename.endswith(".json"):
            gj = json.loads(raw.decode("utf-8"))
            features = gj.get("features", gj if isinstance(gj, list) else [])
            for i, feat in enumerate(features):
                props = feat.get("properties", {}) if isinstance(feat, dict) else {}
                geom = feat.get("geometry", {}) if isinstance(feat, dict) else {}
                coords = geom.get("coordinates") if geom else None
                row = {
                    "external_id": props.get("external_id") or props.get("id") or props.get("camera_id"),
                    "name": props.get("name") or props.get("external_id"),
                    "department_code": props.get("department") or props.get("department_code"),
                    "heading_deg": props.get("heading_deg") or props.get("heading"),
                    "fov_deg": props.get("fov_deg") or props.get("fov"),
                    "range_m": props.get("range_m") or props.get("range"),
                    "rtsp_url": props.get("rtsp_url") or props.get("rtsp"),
                }
                if coords and len(coords) >= 2:
                    row["longitude"], row["latitude"] = float(coords[0]), float(coords[1])
                else:
                    row["latitude"] = props.get("latitude") or props.get("lat")
                    row["longitude"] = props.get("longitude") or props.get("lon")
                rows.append({"_row": i + 1, **row})
        elif filename.endswith(".csv") or filename.endswith(".txt"):
            text_data = raw.decode("utf-8-sig")
            reader = csv.DictReader(io.StringIO(text_data))
            for i, r in enumerate(reader):
                rows.append(
                    {
                        "_row": i + 2,
                        "external_id": r.get("external_id") or r.get("id") or r.get("camera_id"),
                        "name": r.get("name") or r.get("external_id"),
                        "department_code": r.get("department_code") or r.get("department"),
                        "latitude": r.get("latitude") or r.get("lat"),
                        "longitude": r.get("longitude") or r.get("lon"),
                        "heading_deg": r.get("heading_deg") or r.get("heading"),
                        "fov_deg": r.get("fov_deg") or r.get("fov"),
                        "range_m": r.get("range_m") or r.get("range"),
                        "rtsp_url": r.get("rtsp_url") or r.get("rtsp"),
                    }
                )
        elif filename.endswith(".xlsx") or filename.endswith(".xls"):
            try:
                import openpyxl
            except ImportError as e:
                raise HTTPException(400, "openpyxl required for Excel uploads") from e
            wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True)
            ws = wb.active
            headers = [str(c.value).strip().lower() if c.value else "" for c in next(ws.iter_rows(max_row=1))]
            for i, excel_row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
                r = {headers[j]: excel_row[j] for j in range(len(headers)) if headers[j]}
                rows.append(
                    {
                        "_row": i + 2,
                        "external_id": r.get("external_id") or r.get("id") or r.get("camera_id"),
                        "name": r.get("name") or r.get("external_id"),
                        "department_code": r.get("department_code") or r.get("department"),
                        "latitude": r.get("latitude") or r.get("lat"),
                        "longitude": r.get("longitude") or r.get("lon"),
                        "heading_deg": r.get("heading_deg") or r.get("heading"),
                        "fov_deg": r.get("fov_deg") or r.get("fov"),
                        "range_m": r.get("range_m") or r.get("range"),
                        "rtsp_url": r.get("rtsp_url") or r.get("rtsp"),
                    }
                )
        else:
            raise HTTPException(400, "Unsupported format. Use CSV, Excel, or GeoJSON.")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(400, f"Failed to parse upload: {e}") from e

    created = updated = 0
    for row in rows:
        row_num = row.pop("_row", None)
        try:
            if not row.get("external_id") or not row.get("name"):
                raise ValueError("external_id and name are required")
            # OPERATOR: force / validate department scope
            if u.role.value == "OPERATOR" and u.department_code:
                row["department_code"] = u.department_code
            assert_operator_department(u, row.get("department_code"))
            for key in ("latitude", "longitude", "heading_deg", "fov_deg", "range_m"):
                if row.get(key) is not None and row[key] != "":
                    row[key] = float(row[key])
                else:
                    row[key] = None
            payload = CameraCreate(
                external_id=str(row["external_id"]),
                name=str(row["name"]),
                department_code=str(row["department_code"]) if row.get("department_code") else None,
                latitude=row.get("latitude"),
                longitude=row.get("longitude"),
                heading_deg=row.get("heading_deg"),
                fov_deg=row.get("fov_deg"),
                range_m=row.get("range_m"),
                rtsp_url=str(row["rtsp_url"]) if row.get("rtsp_url") else None,
            )
            existing = (
                await db.execute(select(Camera).where(Camera.external_id == payload.external_id))
            ).scalar_one_or_none()
            dept = await _get_or_create_department(db, payload.department_code)
            if existing:
                existing.name = payload.name
                existing.department_id = dept.id if dept else existing.department_id
                existing.latitude = payload.latitude
                existing.longitude = payload.longitude
                existing.heading_deg = payload.heading_deg
                existing.fov_deg = payload.fov_deg
                existing.range_m = payload.range_m
                if payload.rtsp_url:
                    existing.rtsp_url = payload.rtsp_url
                _apply_geometry(existing)
                updated += 1
            else:
                cam = Camera(
                    external_id=payload.external_id,
                    name=payload.name,
                    department_id=dept.id if dept else None,
                    latitude=payload.latitude,
                    longitude=payload.longitude,
                    heading_deg=payload.heading_deg,
                    fov_deg=payload.fov_deg,
                    range_m=payload.range_m,
                    rtsp_url=payload.rtsp_url,
                    whep_path=payload.external_id,
                    hls_path=payload.external_id,
                    status=CameraStatus.UNKNOWN,
                )
                _apply_geometry(cam)
                db.add(cam)
                created += 1
        except Exception as e:
            errors.append({"row": row_num, "error": str(e), "data": row})

    await db.flush()
    return BulkUploadResult(total=len(rows), created=created, updated=updated, errors=errors)
